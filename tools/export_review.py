#!/usr/bin/env python3
"""Export items whose structural fields could not be read from the client text
to an Excel workbook for manual review, and read the corrections back.

  python3 tools/export_review.py            → data/review_items.xlsx
  python3 tools/export_review.py --import   → data/manual_overrides.json (used by convert_items.py)

Sheets: weapons (subType / weaponLevel / atk), armor (def / jobs). Fill the
"FIX_*" columns and leave the others alone.
"""
import json, sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
ITEMS = ROOT / 'data' / 'items.json'
XLSX = ROOT / 'data' / 'review_items.xlsx'
OVERRIDES = ROOT / 'data' / 'manual_overrides.json'

WEAPON_TYPES = ['DAGGER', 'SWORD_1H', 'SWORD_2H', 'SPEAR_1H', 'SPEAR_2H', 'AXE_1H', 'AXE_2H', 'MACE', 'STAFF_1H', 'STAFF_2H',
                'BOW', 'KATAR', 'BOOK', 'KNUCKLE', 'INSTRUMENT', 'WHIP', 'HUUMA', 'REVOLVER', 'RIFLE', 'SHOTGUN', 'GATLING', 'GRENADE_LAUNCHER']

def sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        if c.value.startswith('FIX_'):
            c.fill = PatternFill('solid', fgColor='FFF2CC')
    for r in rows: ws.append(r)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = 60 if h == 'description' else 22
    ws.freeze_panes = 'A2'
    return ws

def export():
    items = json.load(open(ITEMS, encoding='utf-8'))
    wb = Workbook(); wb.remove(wb.active)
    weapons = [i for i in items if i['itemType'] == 'WEAPON' and (i['subType'] == 'UNKNOWN' or i['weaponLevel'] is None or i['atk'] is None)]
    sheet(wb, 'weapons',
          ['id', 'name', 'subType(now)', 'weaponLevel(now)', 'atk(now)', 'FIX_subType', 'FIX_weaponLevel', 'FIX_atk', 'description'],
          [[i['id'], i['name'], i['subType'], i['weaponLevel'], i['atk'], '', '', '', i['description']] for i in weapons])
    armor = [i for i in items if i['itemType'] == 'ARMOR' and (i['def'] is None or not i['jobs'])]
    sheet(wb, 'armor',
          ['id', 'name', 'subType', 'def(now)', 'jobs(now)', 'FIX_def', 'FIX_jobs', 'description'],
          [[i['id'], i['name'], i['subType'], i['def'], i['jobs'], '', '', i['description']] for i in armor])
    ws = wb.create_sheet('README')
    ws.append(['FIX_subType values:', ', '.join(WEAPON_TYPES)])
    ws.append(['FIX_weaponLevel:', '1-4'])
    ws.append(['FIX_atk / FIX_def:', 'number'])
    ws.append(['FIX_jobs:', 'free text, e.g. "Swordman Cls/Merchant Cls"'])
    ws.append(['Then run:', 'python3 tools/export_review.py --import && python3 tools/convert_items.py && (cd server && bun run db:seed)'])
    wb.save(XLSX)
    print(f'weapons {len(weapons)}, armor {len(armor)} → {XLSX}')

def import_fixes():
    wb = load_workbook(XLSX)
    overrides = json.load(open(OVERRIDES, encoding='utf-8')) if OVERRIDES.exists() else {}
    n = 0
    for name in ('weapons', 'armor'):
        ws = wb[name]
        headers = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(headers, row))
            fix = {h[4:]: v for h, v in rec.items() if h and h.startswith('FIX_') and v not in (None, '')}
            if not fix: continue
            for k in ('weaponLevel', 'atk', 'def'):
                if k in fix: fix[k] = int(fix[k])
            if 'subType' in fix: fix['subType'] = str(fix['subType']).strip().upper()
            overrides.setdefault(str(rec['id']), {}).update(fix); n += 1
    json.dump(overrides, open(OVERRIDES, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'{n} rows imported → {OVERRIDES} ({len(overrides)} items total)')

if __name__ == '__main__':
    import_fixes() if '--import' in sys.argv else export()
